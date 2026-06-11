"""Generate a PPTX describing the latest Health Platform capabilities.

This script intentionally keeps dependencies minimal and produces a self-contained deck.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


# 16:9 widescreen in inches
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


class Palette:
    bg = RGBColor(250, 251, 252)  # near-white
    text = RGBColor(27, 31, 35)  # near-black
    muted = RGBColor(88, 96, 105)
    primary = RGBColor(39, 120, 132)  # deep teal
    accent = RGBColor(254, 68, 71)  # coral/red
    card = RGBColor(255, 255, 255)
    border = RGBColor(219, 224, 229)


def _set_font(run, size_pt: int, bold: bool = False, color: RGBColor | None = None, name: str = "Arial"):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def add_bg(slide, color: RGBColor):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, title: str, subtitle: str | None = None):
    add_bg(slide, Palette.bg)

    title_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.2), Inches(11.6), Inches(1.2))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = title
    _set_font(r, 44, bold=True, color=Palette.primary)

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.9), Inches(2.2), Inches(11.6), Inches(0.8))
        stf = sub_box.text_frame
        stf.clear()
        sp = stf.paragraphs[0]
        sp.alignment = PP_ALIGN.LEFT
        sr = sp.add_run()
        sr.text = subtitle
        _set_font(sr, 18, color=Palette.muted)

    # Accent bar
    slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.9),
        Inches(0.95),
        Inches(1.8),
        Inches(0.14),
    ).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = Palette.accent
    slide.shapes[-1].line.fill.background()


def add_section_header(slide, title: str, tag: str | None = None):
    add_bg(slide, Palette.bg)

    # Header area
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.6),
        Inches(0.5),
        Inches(12.1),
        Inches(0.9),
    )
    header.fill.solid()
    header.fill.fore_color.rgb = Palette.card
    header.line.color.rgb = Palette.border

    tb = slide.shapes.add_textbox(Inches(0.9), Inches(0.66), Inches(10.5), Inches(0.6))
    tf = tb.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = title
    _set_font(r, 28, bold=True, color=Palette.text)

    if tag:
        chip = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(11.2),
            Inches(0.67),
            Inches(1.2),
            Inches(0.45),
        )
        chip.fill.solid()
        chip.fill.fore_color.rgb = Palette.primary
        chip.line.fill.background()
        ctf = chip.text_frame
        ctf.clear()
        cp = ctf.paragraphs[0]
        cp.alignment = PP_ALIGN.CENTER
        cr = cp.add_run()
        cr.text = tag
        _set_font(cr, 12, bold=True, color=RGBColor(255, 255, 255))


def add_bullets(slide, left: float, top: float, width: float, height: float, items: list[str]):
    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    box.fill.solid()
    box.fill.fore_color.rgb = Palette.card
    box.line.color.rgb = Palette.border

    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True

    for idx, text in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(6)
        p.text = text
        p.font.name = "Arial"
        p.font.size = Pt(18)
        p.font.color.rgb = Palette.text


def add_two_columns(slide, left_items: list[str], right_items: list[str], left_title: str, right_title: str):
    # Left column title
    lt = slide.shapes.add_textbox(Inches(0.9), Inches(1.6), Inches(5.9), Inches(0.4))
    ltf = lt.text_frame
    ltf.clear()
    lp = ltf.paragraphs[0]
    lr = lp.add_run()
    lr.text = left_title
    _set_font(lr, 18, bold=True, color=Palette.primary)

    # Right column title
    rt = slide.shapes.add_textbox(Inches(7.0), Inches(1.6), Inches(5.9), Inches(0.4))
    rtf = rt.text_frame
    rtf.clear()
    rp = rtf.paragraphs[0]
    rr = rp.add_run()
    rr.text = right_title
    _set_font(rr, 18, bold=True, color=Palette.primary)

    add_bullets(slide, 0.9, 2.05, 5.9, 4.9, left_items)
    add_bullets(slide, 7.0, 2.05, 5.9, 4.9, right_items)


def add_architecture_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "平台架构与分层", tag="ARCH")

    # 4-layer diagram
    x0, y0 = Inches(0.9), Inches(2.0)
    w, h = Inches(11.6), Inches(0.9)
    layers = [
        ("Client", "React 18 + Ant Design + ECharts + i18next", Palette.primary),
        ("Service", "Flask API：参数校验 / 认证鉴权 / HTTP 语义", Palette.accent),
        ("Manager", "业务逻辑 + DB 操作（禁止 Service 直连 DB）", Palette.primary),
        ("Models", "SQLAlchemy Models + Migrations", Palette.accent),
    ]

    for i, (name, desc, color) in enumerate(layers):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x0, y0 + Inches(i * 1.05), w, h)
        box.fill.solid()
        box.fill.fore_color.rgb = Palette.card
        box.line.color.rgb = Palette.border

        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x0, y0 + Inches(i * 1.05), Inches(1.3), h)
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()

        t1 = slide.shapes.add_textbox(x0 + Inches(0.18), y0 + Inches(i * 1.05) + Inches(0.22), Inches(1.1), Inches(0.5))
        tf1 = t1.text_frame
        tf1.clear()
        p1 = tf1.paragraphs[0]
        p1.alignment = PP_ALIGN.LEFT
        r1 = p1.add_run()
        r1.text = name
        _set_font(r1, 16, bold=True, color=RGBColor(255, 255, 255))

        t2 = slide.shapes.add_textbox(x0 + Inches(1.55), y0 + Inches(i * 1.05) + Inches(0.22), Inches(9.8), Inches(0.6))
        tf2 = t2.text_frame
        tf2.clear()
        p2 = tf2.paragraphs[0]
        r2 = p2.add_run()
        r2.text = desc
        _set_font(r2, 18, color=Palette.text)


def build_deck(out_path: Path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    today = datetime.now().strftime("%Y-%m-%d")

    # Slide 1: Cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(
        slide,
        "Health Platform｜产品最新能力概览",
        f"{today}  ·  AKS Shared Cluster + GHCR + Agent Skills",
    )

    # Slide 2: Executive summary
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "一句话总结", tag="TL;DR")
    add_bullets(
        slide,
        0.9,
        1.8,
        11.6,
        5.6,
        [
            "从个人健康记录应用升级为可演示/可训练/可扩展的工程化平台：应用 + DevOps + 安全 + 代理技能。",
            "支持共享 AKS 多租户（branch → namespace）隔离部署，前端 LoadBalancer 暴露，后端 ClusterIP 内网访问。",
            "引入 GitHub Advanced Security 与 Copilot 工作流，覆盖：依赖漏洞、secret scanning、PR 自动审查与修复建议。",
        ],
    )

    # Slide 3: Product capabilities
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "产品核心能力", tag="APP")
    add_two_columns(
        slide,
        left_title="健康记录与家庭成员",
        right_title="用户权限与可视化",
        left_items=[
            "血压（收缩压/舒张压）、心率、标签、备注记录",
            "数据校验（血压 30–250；收缩压 > 舒张压；心率 30–150）",
            "分页查询（page + size），按日期/标签筛选",
            "家庭成员管理：Self/自己保护（不可编辑/删除）",
            "CSV 导出（UTF-8 BOM + RFC5987 filename*）",
        ],
        right_items=[
            "JWT 登录体系：注册/登录/刷新/登出",
            "角色：USER / ADMIN / SUPER_ADMIN",
            "管理员重置密码：临时密码 + must_change_password + token_version",
            "趋势图：周/月/全部，异常高亮，PNG 下载",
            "i18next 多语言与 Ant Design 组件化 UI",
        ],
    )

    # Slide 4: Architecture
    add_architecture_slide(prs)

    # Slide 5: Deployment model
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "部署与多租户模型（共享 AKS）", tag="AKS")
    add_bullets(
        slide,
        0.9,
        1.7,
        11.6,
        5.8,
        [
            "镜像仓库：GHCR（push 用 GITHUB_TOKEN；pull 用长期 GHCR_READ_TOKEN 创建 imagePullSecret）",
            "分支隔离：CI 根据分支名计算 namespace（ws-<normalized>-<hash>）并自动创建",
            "安全护栏：namespace 级 ResourceQuota + LimitRange（pods/services/loadbalancers/CPU/Memory）",
            "服务暴露：frontend-svc 可配置为 LoadBalancer；backend-svc 默认为 ClusterIP",
            "运行安全：容器 non-root，最小权限 capabilities，readiness/liveness 探针",
        ],
    )

    # Slide 6: CI/CD pipeline highlights
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "CI/CD 工作流亮点", tag="CI")
    add_two_columns(
        slide,
        left_title="Build & Push（GHCR）",
        right_title="Deploy（Kubernetes）",
        left_items=[
            "按变更判断是否需要 build（代码/依赖/前端变更触发）",
            "多 tag 策略：branch / sha / latest(main) / mvp(MVP*)",
            "强制 tag 小写，避免 registry 不一致",
            "Docker Buildx + metadata-action + build-push-action",
        ],
        right_items=[
            "环境分层：GitHub Secrets > GitHub Vars > deploy/config/*.env",
            "模板渲染：envsubst 生成 deploy/k8s-<env>-generated.yaml",
            "自动创建 ghcr-secret（docker-registry secret）",
            "rollout status 失败自动采集诊断（pods/events/describe/logs）",
        ],
    )

    # Slide 7: Security & governance
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "安全与治理", tag="SEC")
    add_bullets(
        slide,
        0.9,
        1.7,
        11.6,
        5.8,
        [
            "应用侧：JWT 鉴权、CORS 白名单、Flask-Limiter 限流、输入校验与错误语义统一",
            "平台侧：GitHub Advanced Security（secret scanning push protection / Dependabot / Code Quality）",
            "Copilot 参与 PR：自动审查、建议修复依赖漏洞/配置问题",
            "敏感信息策略：.env 不入库；生产密钥走 GitHub Secrets；K8s 侧不打印 secret 内容",
        ],
    )

    # Slide 8: Observability & resilience
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "可观测与可靠性", tag="SRE")
    add_two_columns(
        slide,
        left_title="运行可观测",
        right_title="可靠性能力",
        left_items=[
            "/api/healthz 健康检查 + K8s 探针驱动自愈",
            "kubectl logs/describe/events 快速定位故障",
            "CI 失败自动采集诊断输出，减少 MTTR",
        ],
        right_items=[
            "tenacity 重试策略（指数退避）",
            "pybreaker 熔断（fail_max/reset_timeout）",
            "资源隔离：namespace 配额避免噪声邻居问题",
        ],
    )

    # Slide 9: Agent Skills
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "Agent Skills 能力矩阵（新增）", tag="AI")
    add_two_columns(
        slide,
        left_title="Document Skills",
        right_title="Engineering & Workflow Skills",
        left_items=[
            "pptx：演示文稿创建/编辑/分析（OOXML + html2pptx 流程）",
            "docx：Word 文档生成/编辑",
            "xlsx：表格生成/计算/重算",
            "pdf：提取/表单处理/分析",
        ],
        right_items=[
            "webapp-testing：端到端网页测试方法与脚本化流程",
            "mcp-builder：快速生成 MCP server 参考实现与评估",
            "internal-comms：3P/公告/FAQ 等企业沟通模板",
            "theme-factory / brand-guidelines / frontend-design：设计与一致性产出",
            "web-artifacts-builder / skill-creator / slack-gif-creator：内容与技能扩展",
        ],
    )

    # Slide 10: Typical scenarios
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "典型应用场景", tag="USE")
    add_bullets(
        slide,
        0.9,
        1.7,
        11.6,
        5.8,
        [
            "面向管理层：一键生成周报/里程碑汇报 PPT（internal-comms + pptx）",
            "面向研发：自动化排障与命令生成（Copilot + kubectl + Troubleshooting playbook）",
            "面向安全：演示 secret scanning 与依赖漏洞闭环（GHAS + Dependabot + PR Review）",
            "面向平台：共享集群多租户实践（branch→namespace、quota、pull secret、LB）",
        ],
    )

    # Slide 11: Roadmap
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "后续增强方向（建议）", tag="NEXT")
    add_bullets(
        slide,
        0.9,
        1.7,
        11.6,
        5.8,
        [
            "可观测：接入标准化指标与日志聚合（Prometheus/Grafana/Log pipeline）",
            "安全：更严格的 Policy-as-Code（OPA/Gatekeeper）与镜像签名（cosign）",
            "部署：后端/前端统一 ingress + TLS，减少 LoadBalancer 消耗",
            "技能：为 Health Platform 定制“产品发布包”技能（PPT + Release notes + FAQ）",
        ],
    )

    # Slide 12: Closing
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "Thanks", tag="END")
    add_bullets(
        slide,
        0.9,
        1.9,
        11.6,
        5.4,
        [
            "仓库：健康记录平台（Flask + React）",
            "交付：共享 AKS 部署流程 + GHAS 安全练习 + Agent Skills 扩展",
            "建议下一步：用模板化流程把演示资产（PPT/Doc/Runbook）纳入 CI 产物",
        ],
    )

    prs.save(out_path)


def main():
    repo_root = Path(__file__).resolve().parents[1]
    out_dir = repo_root / "docs" / "reports"
    out_dir.mkdir(parents=True, exist_ok=True)

    out_path = out_dir / "health_platform_latest_features_2026-01-08.pptx"
    build_deck(out_path)
    print(str(out_path))


if __name__ == "__main__":
    main()
