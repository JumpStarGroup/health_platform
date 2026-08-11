"""Generate a PPTX deck for the Health Platform workflow onboarding training.

This script uses python-pptx and keeps the deck self-contained.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


class Palette:
    bg = RGBColor(246, 248, 250)
    text = RGBColor(28, 33, 39)
    muted = RGBColor(93, 103, 115)
    primary = RGBColor(24, 88, 124)
    accent = RGBColor(239, 115, 46)
    success = RGBColor(43, 138, 64)
    danger = RGBColor(194, 58, 58)
    card = RGBColor(255, 255, 255)
    border = RGBColor(220, 226, 233)


def set_font(run, size_pt: int, bold: bool = False, color: RGBColor | None = None, name: str = "Microsoft YaHei"):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def fill_background(slide, color: RGBColor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header(slide, title: str, tag: str | None = None):
    fill_background(slide, Palette.bg)

    header = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(0.45), Inches(12.15), Inches(0.95))
    header.fill.solid()
    header.fill.fore_color.rgb = Palette.card
    header.line.color.rgb = Palette.border

    title_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.66), Inches(10.5), Inches(0.45))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_font(r, 26, bold=True, color=Palette.primary)

    if tag:
        chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.35), Inches(0.64), Inches(1.15), Inches(0.42))
        chip.fill.solid()
        chip.fill.fore_color.rgb = Palette.primary
        chip.line.fill.background()
        tf = chip.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = tag
        set_font(r, 11, bold=True, color=RGBColor(255, 255, 255))


def add_title_slide(slide, title: str, subtitle: str, footer: str):
    fill_background(slide, Palette.bg)

    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.25)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = Palette.primary
    slide.shapes[-1].line.fill.background()

    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.2), Inches(2.1), Inches(0.16))
    accent.fill.solid()
    accent.fill.fore_color.rgb = Palette.accent
    accent.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.55), Inches(11.5), Inches(1.2))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_font(r, 34, bold=True, color=Palette.primary)

    subtitle_box = slide.shapes.add_textbox(Inches(0.9), Inches(2.7), Inches(11.6), Inches(0.65))
    tf = subtitle_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = subtitle
    set_font(r, 18, color=Palette.muted)

    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(3.55), Inches(4.8), Inches(0.55))
    badge.fill.solid()
    badge.fill.fore_color.rgb = Palette.card
    badge.line.color.rgb = Palette.border
    tf = badge.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = footer
    set_font(r, 13, bold=True, color=Palette.primary)

    for idx, x in enumerate([8.8, 9.65, 10.5, 11.35, 12.2]):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(4.15), Inches(0.24), Inches(0.24))
        dot.fill.solid()
        dot.fill.fore_color.rgb = Palette.accent if idx % 2 == 0 else Palette.success
        dot.line.fill.background()


def add_bullet_card(slide, left: float, top: float, width: float, height: float, title: str, items: list[str], tag_color: RGBColor | None = None):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = Palette.card
    card.line.color.rgb = Palette.border

    title_box = slide.shapes.add_textbox(Inches(left + 0.18), Inches(top + 0.12), Inches(width - 0.36), Inches(0.32))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_font(r, 16, bold=True, color=tag_color or Palette.primary)

    body = slide.shapes.add_textbox(Inches(left + 0.18), Inches(top + 0.5), Inches(width - 0.36), Inches(height - 0.62))
    tf = body.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, text in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(5)
        p.text = text
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(16)
        p.font.color.rgb = Palette.text


def add_flow_boxes(slide, boxes: list[tuple[str, str, RGBColor]], y: float = 2.05):
    x = 0.7
    width = 2.15
    gap = 0.18
    for idx, (title, desc, color) in enumerate(boxes):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + idx * (width + gap)), Inches(y), Inches(width), Inches(1.15))
        box.fill.solid()
        box.fill.fore_color.rgb = color
        box.line.fill.background()

        title_box = slide.shapes.add_textbox(Inches(x + 0.1 + idx * (width + gap)), Inches(y + 0.14), Inches(width - 0.2), Inches(0.3))
        tf = title_box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = title
        set_font(r, 14, bold=True, color=RGBColor(255, 255, 255))

        desc_box = slide.shapes.add_textbox(Inches(x + 0.1 + idx * (width + gap)), Inches(y + 0.48), Inches(width - 0.2), Inches(0.48))
        tf = desc_box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = desc
        set_font(r, 11, color=RGBColor(255, 255, 255))

        if idx < len(boxes) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x + (idx + 1) * (width + gap) - 0.15), Inches(y + 0.36), Inches(0.28), Inches(0.42))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = Palette.accent
            arrow.line.fill.background()


def build_deck(out_path: Path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    version = out_path.stem.split("_v")[-1]
    today = datetime.now().strftime("%Y-%m-%d")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_slide(
        slide,
        "Health Platform 研发流程与角色分工",
        "新员工快速上手培训版：从需求到上线的协作规范",
        f"{today}  ·  Version {version}",
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "为什么需要这套流程", "WHY")
    add_bullet_card(
        slide,
        0.75,
        1.55,
        5.95,
        4.9,
        "统一流程的价值",
        [
            "保证需求可追踪，Issue 是跨角色主线。",
            "让 main 始终保持可部署状态。",
            "让文档、代码、测试、发布彼此回链。",
            "减少返工，降低多人协作沟通成本。",
        ],
    )
    add_bullet_card(
        slide,
        6.95,
        1.55,
        5.65,
        4.9,
        "新人最常见的问题",
        [
            "什么时候建 feature 分支？",
            "docs PR 和实现 PR 的区别是什么？",
            "Issue 应该什么时候关闭？",
            "上线前需要经过哪些门禁？",
        ],
        tag_color=Palette.accent,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "端到端流程总览", "FLOW")
    add_flow_boxes(
        slide,
        [
            ("PM", "澄清需求\n创建 Issue", Palette.primary),
            ("审批", "审范围\n控风险", Palette.accent),
            ("Architect", "输出 design\n定义方案", Palette.primary),
            ("Planner", "拆 plan\n定验证", Palette.accent),
            ("Dev", "编码\n测试\n提交 PR", Palette.primary),
            ("Release", "版本\ntag\n上线", Palette.success),
        ],
        y=2.15,
    )
    add_bullet_card(
        slide,
        0.85,
        4.0,
        11.55,
        1.75,
        "一句话总结",
        ["从想法到上线，每一步都有明确角色、产物和门禁。"],
        tag_color=Palette.success,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "角色与职责", "ROLE")
    add_bullet_card(
        slide,
        0.75,
        1.55,
        6.0,
        4.95,
        "前半程：定义与决策",
        [
            "Product_Manager：定义问题、范围、价值。",
            "需求审批：检查目标、风险、验收标准。",
            "System_Architect：定义系统方案与边界。",
            "Tech_Lead_Planner：把方案拆成可执行计划。",
        ],
    )
    add_bullet_card(
        slide,
        6.95,
        1.55,
        5.65,
        4.95,
        "后半程：实现与交付",
        [
            "Developer：编码、测试、联调、提交 PR。",
            "QA / Playwright：维护 E2E 和关键路径回归。",
            "Reviewer：把关代码质量与可维护性。",
            "Release_Manager：版本、CHANGELOG、发布说明。",
        ],
        tag_color=Palette.accent,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "需求阶段：先讲清楚问题", "REQ")
    add_bullet_card(
        slide,
        0.75,
        1.55,
        6.1,
        4.95,
        "Product_Manager 的动作",
        [
            "先确认问题是什么，再谈怎么做。",
            "创建或关联 GitHub Issue。",
            "输出需求文档，说明 What / Why / 范围 / 验收。",
            "不讨论数据库、API、实现细节。",
        ],
    )
    add_bullet_card(
        slide,
        7.0,
        1.55,
        5.6,
        4.95,
        "分支与 PR 规则",
        [
            "需求/设计/计划阶段统一使用 docs/<issue>-<slug>。",
            "docs PR 只用 Refs #<issue>。",
            "不要在需求阶段直接创建 feature 分支。",
        ],
        tag_color=Palette.success,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "什么时候创建 Feature 分支", "BRANCH")
    add_bullet_card(
        slide,
        0.75,
        1.55,
        5.9,
        4.95,
        "推荐时机",
        [
            "需求、设计、计划确认后，再进入实现。",
            "从最新 main 创建 feature/* 或 fix/*。",
            "这样实现分支天然包含已批准文档。",
        ],
    )
    add_bullet_card(
        slide,
        6.95,
        1.55,
        5.65,
        4.95,
        "流程图",
        ["docs 分支完成文档 → docs PR 合入 main → Developer 从最新 main 拉出 feature/fix → 开始编码。"],
        tag_color=Palette.accent,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Developer 的完整闭环", "DEV")
    add_bullet_card(
        slide,
        0.75,
        1.55,
        5.95,
        4.95,
        "标准动作",
        [
            "读取 plan / requirement / design。",
            "先写测试，再写代码。",
            "本地验证后提交 Conventional Commits。",
            "push 分支并发起 PR。",
        ],
    )
    add_bullet_card(
        slide,
        6.95,
        1.55,
        5.65,
        4.95,
        "本地三终端模型",
        [
            "Terminal 1：后端服务。",
            "Terminal 2：前端开发服务器。",
            "Terminal 3：测试、Git、一次性命令。",
            "合并后清理分支，保持工作区干净。",
        ],
        tag_color=Palette.primary,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Issue 与 PR 的关系", "PR")
    add_bullet_card(
        slide,
        0.75,
        1.55,
        12.0,
        2.05,
        "统一原则",
        ["Issue 是主线，PR 是载体。文档先行，代码后跟；只有真正完成验收的实现 PR 才适合关闭 Issue。"],
        tag_color=Palette.success,
    )
    table = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(3.85), Inches(12.0), Inches(2.2))
    table.fill.solid()
    table.fill.fore_color.rgb = Palette.card
    table.line.color.rgb = Palette.border
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(4.1), Inches(11.5), Inches(1.8))
    tf = tb.text_frame
    tf.clear()
    lines = [
        "docs PR：Refs #123，不关闭 Issue",
        "feature / fix PR：Closes #123 或 Fixes #123 视情况使用",
        "release PR：Refs #123，不代表需求完成",
        "follow-up PR：Refs #123，保留追踪关系",
    ]
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(17)
        p.font.color.rgb = Palette.text
        p.space_after = Pt(4)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "测试与质量门禁", "QA")
    add_bullet_card(
        slide,
        0.75,
        1.55,
        5.95,
        4.95,
        "必须经过的验证",
        [
            "单元测试：Pytest。",
            "UI 回归：Playwright E2E。",
            "PR 校验：后端测试 + 前端 build。",
            "发布前再检查版本文件和发布说明。",
        ],
    )
    add_bullet_card(
        slide,
        6.95,
        1.55,
        5.65,
        4.95,
        "质量门禁的意义",
        [
            "任何一步不过，都不能继续往下走。",
            "门禁不是阻碍交付，而是避免把问题扩散到生产。",
        ],
        tag_color=Palette.danger,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "分支策略与发布", "GIT")
    add_bullet_card(
        slide,
        0.75,
        1.55,
        6.0,
        4.95,
        "分支类型",
        [
            "main：稳定主干。",
            "docs/<issue>-<slug>：文档协作。",
            "feature/<scope>-<desc>：新功能开发。",
            "fix/<scope>-<desc>：缺陷修复。",
            "release/<version> / hotfix/<version>：发布与紧急修复。",
        ],
    )
    add_bullet_card(
        slide,
        6.95,
        1.55,
        5.65,
        4.95,
        "发布节奏",
        [
            "合并到 main 后部署 staging。",
            "staging 回归通过后创建 release。",
            "tag 驱动生产发布。",
            "版本、CHANGELOG、Release Notes 必须一致。",
        ],
        tag_color=Palette.success,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "新员工上手清单", "CHECK")
    add_bullet_card(
        slide,
        0.75,
        1.55,
        6.0,
        4.95,
        "入职第一周建议完成",
        [
            "看懂流程总览图。",
            "理解角色分工和责任边界。",
            "理解 docs / feature / fix / release 的区别。",
            "理解 Issue 和 PR 的关联规则。",
        ],
    )
    add_bullet_card(
        slide,
        6.95,
        1.55,
        5.65,
        4.95,
        "实践目标",
        [
            "能独立启动本地环境。",
            "能跟着一个真实 Issue 跑到 PR。",
            "能在 review 和测试反馈后完成闭环。",
        ],
        tag_color=Palette.primary,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Q&A", "END")
    add_bullet_card(
        slide,
        1.0,
        1.95,
        11.25,
        3.45,
        "最后一句话",
        [
            "需求先行，Issue 贯穿，docs 先于 feature，feature 先于 release，质量门禁贯穿全程。",
            "建议下一步：跟一个真实需求从 Issue 跑到 PR。",
        ],
        tag_color=Palette.success,
    )

    prs.save(out_path)


def main():
    repo_root = Path(__file__).resolve().parents[1]
    version = (repo_root / "VERSION").read_text(encoding="utf-8").strip()
    today = datetime.now().strftime("%Y-%m-%d")
    out_path = repo_root / "docs" / "reports" / f"health_platform_workflow_onboarding_{today}_v{version}.pptx"
    build_deck(out_path)
    print(out_path)


if __name__ == "__main__":
    main()