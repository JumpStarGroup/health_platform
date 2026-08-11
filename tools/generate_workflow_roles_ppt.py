"""Generate a PPTX deck for the Health Platform workflow and role division."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


class Palette:
    bg = RGBColor(242, 245, 248)
    card = RGBColor(255, 255, 255)
    text = RGBColor(28, 33, 39)
    muted = RGBColor(92, 103, 115)
    primary = RGBColor(17, 78, 120)
    accent = RGBColor(236, 118, 46)
    teal = RGBColor(48, 150, 145)
    green = RGBColor(47, 140, 76)
    red = RGBColor(194, 58, 58)
    border = RGBColor(219, 226, 233)


def set_font(run, size: int, *, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def fill_bg(slide) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = Palette.bg


def add_title_bar(slide, title: str, chip: str | None = None) -> None:
    fill_bg(slide)
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(0.38), Inches(12.2), Inches(0.88))
    bar.fill.solid()
    bar.fill.fore_color.rgb = Palette.card
    bar.line.color.rgb = Palette.border

    title_box = slide.shapes.add_textbox(Inches(0.88), Inches(0.58), Inches(10.5), Inches(0.4))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_font(r, 25, bold=True, color=Palette.primary)

    if chip:
        chip_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.1), Inches(0.57), Inches(1.35), Inches(0.42))
        chip_box.fill.solid()
        chip_box.fill.fore_color.rgb = Palette.primary
        chip_box.line.fill.background()
        tf = chip_box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = chip
        set_font(r, 11, bold=True, color=RGBColor(255, 255, 255))


def add_textbox(slide, left: float, top: float, width: float, height: float, text: str, *, size: int = 16, color: RGBColor = Palette.text, bold: bool = False, align: PP_ALIGN = PP_ALIGN.LEFT) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    set_font(r, size, bold=bold, color=color)


def add_card(slide, left: float, top: float, width: float, height: float, title: str, items: list[str], *, header_color: RGBColor = Palette.primary) -> None:
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = Palette.card
    card.line.color.rgb = Palette.border

    add_textbox(slide, left + 0.2, top + 0.15, width - 0.4, 0.28, title, size=16, color=header_color, bold=True)

    body = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.5), Inches(width - 0.4), Inches(height - 0.62))
    tf = body.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(4)
        p.text = item
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(15)
        p.font.color.rgb = Palette.text


def add_arrow(slide, left: float, top: float, width: float = 0.32) -> None:
    arrow = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(left), Inches(top), Inches(width), Inches(0.36))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = Palette.accent
    arrow.line.fill.background()


def add_stage_box(slide, left: float, top: float, width: float, title: str, desc: str, color: RGBColor) -> None:
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(1.18))
    box.fill.solid()
    box.fill.fore_color.rgb = color
    box.line.fill.background()

    add_textbox(slide, left + 0.08, top + 0.12, width - 0.16, 0.26, title, size=14, color=RGBColor(255, 255, 255), bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, left + 0.08, top + 0.48, width - 0.16, 0.44, desc, size=11, color=RGBColor(255, 255, 255), align=PP_ALIGN.CENTER)


def add_footer(slide, text: str) -> None:
    add_textbox(slide, 0.9, 6.95, 11.6, 0.22, text, size=10, color=Palette.muted)


def build_deck(out_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    today = datetime.now().strftime("%Y-%m-%d")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill_bg(slide)
    head = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.2))
    head.fill.solid()
    head.fill.fore_color.rgb = Palette.primary
    head.line.fill.background()
    add_textbox(slide, 0.9, 1.45, 11.0, 0.7, "Health Platform 研发流程与角色分工", size=32, color=Palette.primary, bold=True)
    add_textbox(slide, 0.92, 2.25, 10.0, 0.45, "从需求到上线的协作规范", size=18, color=Palette.muted)
    add_card(slide, 0.92, 3.05, 5.1, 1.02, "适用对象", ["产品、研发、测试、技术管理、发布协作人员"], header_color=Palette.accent)
    add_card(slide, 6.2, 3.05, 5.1, 1.02, "核心目标", ["让需求、文档、代码、测试、发布全程可追踪"], header_color=Palette.teal)
    add_textbox(slide, 0.92, 4.52, 10.8, 0.32, f"{today}", size=11, color=Palette.muted)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(slide, "为什么需要这套流程", "WHY")
    add_card(slide, 0.75, 1.55, 5.95, 4.95, "统一流程的价值", [
        "需求先定义清楚，再进入设计和开发。",
        "main 始终保持可部署状态。",
        "Issue、文档、PR、测试结果可以回链。",
        "减少多人协作中的返工和沟通成本。",
    ], header_color=Palette.primary)
    add_card(slide, 6.95, 1.55, 5.6, 4.95, "不统一时的典型问题", [
        "需求口头传递，范围不断漂移。",
        "开发直接开工，后面再补文档。",
        "测试点和验收标准不一致。",
        "上线责任、回滚和版本边界不清。",
    ], header_color=Palette.red)
    add_footer(slide, "一句话：需求先行，Issue 贯穿，docs 先于 feature，质量门禁贯穿全程。")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(slide, "端到端流程总览", "FLOW")
    stages = [
        ("PM", "澄清需求\n创建 Issue", Palette.primary),
        ("审批", "审范围\n控风险", Palette.accent),
        ("Architect", "输出 design\n定义方案", Palette.teal),
        ("Planner", "拆 plan\n定验证", Palette.primary),
        ("Dev", "编码\n测试\n提交 PR", Palette.accent),
        ("Release", "版本\ntag\n上线", Palette.green),
    ]
    x = 0.58
    for idx, (title, desc, color) in enumerate(stages):
        add_stage_box(slide, x + idx * 2.12, 2.15, 1.82, title, desc, color)
        if idx < len(stages) - 1:
            add_arrow(slide, x + idx * 2.12 + 1.85, 2.46)
    add_card(slide, 0.88, 4.02, 11.55, 1.55, "流程结论", ["每一步都有明确角色、产物和门禁；任何一步不清楚，都不要直接进入下一步。"], header_color=Palette.green)
    add_footer(slide, "主线顺序：需求澄清 → 需求审批 → 技术设计 → 计划拆解 → 开发测试 → 评审发布。")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(slide, "角色与职责总表", "ROLE")
    rows = [
        ("Product_Manager", "澄清需求、定义价值、控制范围", "需求文档、Issue"),
        ("需求审批", "审范围、风险、验收标准", "评审结论、澄清问题"),
        ("System_Architect", "设计系统方案", "design 文档"),
        ("Tech_Lead_Planner", "拆任务、定阶段、定验证", "plan 文档"),
        ("Developer", "编码、测试、联调、提交 PR", "代码、测试、PR"),
        ("QA / Playwright", "设计和维护 E2E", "测试计划、E2E 用例"),
        ("Reviewer", "代码审查", "Review 结论"),
        ("Release_Manager", "版本发布准备", "VERSION、CHANGELOG、Release Notes"),
        ("DevOps / CI/CD", "构建、部署、回归", "流水线、环境、报告"),
    ]
    table = slide.shapes.add_table(len(rows) + 1, 3, Inches(0.7), Inches(1.55), Inches(11.95), Inches(4.9)).table
    table.columns[0].width = Inches(2.15)
    table.columns[1].width = Inches(5.1)
    table.columns[2].width = Inches(4.7)
    headers = ["角色", "核心职责", "主要产出"]
    for c, header in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = Palette.primary
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                set_font(r, 12, bold=True, color=RGBColor(255, 255, 255))
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = Palette.card if r_idx % 2 else RGBColor(248, 250, 252)
            for p in cell.text_frame.paragraphs:
                if c_idx == 0:
                    p.alignment = PP_ALIGN.CENTER
                for rr in p.runs:
                    set_font(rr, 11, bold=(c_idx == 0), color=Palette.text)
    add_footer(slide, "先看职责边界，再看交付物；边界清楚，协作才不会重叠。")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(slide, "需求阶段怎么协作", "PM")
    add_card(slide, 0.75, 1.55, 5.8, 4.85, "Product_Manager 负责什么", [
        "先确认问题是什么，再谈怎么做。",
        "创建或关联 GitHub Issue。",
        "输出需求文档 req-*.md。",
        "只讲 What / Why / 范围 / 验收。",
    ], header_color=Palette.primary)
    add_card(slide, 6.75, 1.55, 5.8, 4.85, "需求阶段不要做什么", [
        "不要提前讨论数据库和 API 实现细节。",
        "不要绕过审批直接进入开发。",
        "不要让需求文档只剩概念，没有验收标准。",
        "不要在 main 上零散写需求文档。",
    ], header_color=Palette.red)
    add_footer(slide, "建议使用 docs/<issue>-<slug> 作为需求、设计、计划的协作分支。")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(slide, "为什么要先审批需求", "CHECK")
    add_card(slide, 0.8, 1.6, 3.9, 4.9, "审批重点", [
        "目标是否明确。",
        "范围是否清楚。",
        "验收是否可测试。",
        "风险和依赖是否说明。",
        "回滚和发布是否可操作。",
    ], header_color=Palette.accent)
    add_card(slide, 4.9, 1.6, 3.45, 4.9, "处理结果", [
        "通过。",
        "有条件通过。",
        "退回修订。",
    ], header_color=Palette.green)
    add_card(slide, 8.55, 1.6, 3.55, 4.9, "不通过时怎么办", [
        "补充问题说明。",
        "补充风险清单。",
        "补齐验收口径。",
        "重新评审。",
    ], header_color=Palette.red)
    add_footer(slide, "审批不是走流程，而是把不确定性提前暴露。")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(slide, "设计与计划如何交接", "ARCH")
    add_stage_box(slide, 0.92, 2.0, 2.15, "需求文档", "明确 What / Why", Palette.primary)
    add_arrow(slide, 3.18, 2.36)
    add_stage_box(slide, 3.55, 2.0, 2.15, "design", "说明怎么做", Palette.teal)
    add_arrow(slide, 5.81, 2.36)
    add_stage_box(slide, 6.18, 2.0, 2.15, "plan", "拆成任务", Palette.accent)
    add_arrow(slide, 8.44, 2.36)
    add_stage_box(slide, 8.81, 2.0, 2.15, "Developer", "按计划实现", Palette.green)
    add_card(slide, 0.92, 3.75, 11.55, 2.0, "交接原则", [
        "设计阶段看影响面：后端、前端、数据库、测试。",
        "计划阶段看顺序：先数据和模型，再 API，再 UI，最后验证。",
        "这两个阶段仍然建议在 docs 分支里完成，并保持 Issue 关联。",
    ], header_color=Palette.primary)
    add_footer(slide, "一句话：design 负责方案，plan 负责执行顺序。")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(slide, "什么时候创建 feature 分支", "GIT")
    add_card(slide, 0.75, 1.55, 5.95, 4.95, "推荐时机", [
        "需求文档已经具备可验收 AC。",
        "范围和非范围已经明确。",
        "技术方向已经确认可行。",
        "任务拆解已经完成。",
        "docs-only PR 已合入或明确批准。",
    ], header_color=Palette.green)
    add_card(slide, 6.95, 1.55, 5.6, 4.95, "分支建议", [
        "文档阶段：docs/<issue>-<slug>。",
        "开发阶段：feature/<scope>-<desc>。",
        "缺陷修复：fix/<scope>-<desc>。",
        "发版准备：release/<version>。",
    ], header_color=Palette.primary)
    add_footer(slide, "不要在需求刚开始时就开 feature；文档先行，代码后跟。")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(slide, "开发、测试、评审的闭环", "DEV")
    add_card(slide, 0.72, 1.58, 3.0, 4.75, "Developer", [
        "按 plan 实现。",
        "先测试，再编码。",
        "本地验证后提交 PR。",
    ], header_color=Palette.primary)
    add_card(slide, 3.98, 1.58, 2.9, 4.75, "QA", [
        "补齐关键路径。",
        "维护 E2E 用例。",
        "验证回归。",
    ], header_color=Palette.accent)
    add_card(slide, 7.16, 1.58, 2.9, 4.75, "Reviewer", [
        "看行为回归。",
        "看测试是否足够。",
        "看风险和可维护性。",
    ], header_color=Palette.teal)
    add_card(slide, 10.34, 1.58, 2.0, 4.75, "结果", [
        "通过后合并。",
        "不通过则修订。",
    ], header_color=Palette.green)
    add_footer(slide, "质量门禁不是最后一步，而是开发过程的一部分。")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(slide, "发布与环境", "REL")
    add_stage_box(slide, 0.9, 2.0, 2.2, "main", "稳定主干", Palette.primary)
    add_arrow(slide, 3.22, 2.36)
    add_stage_box(slide, 3.58, 2.0, 2.2, "staging", "回归验证", Palette.accent)
    add_arrow(slide, 5.9, 2.36)
    add_stage_box(slide, 6.26, 2.0, 2.2, "release tag", "版本冻结", Palette.teal)
    add_arrow(slide, 8.58, 2.36)
    add_stage_box(slide, 8.94, 2.0, 2.2, "production", "正式上线", Palette.green)
    add_card(slide, 0.95, 3.88, 11.42, 1.92, "发布原则", [
        "main 永远保持可部署。",
        "发布由 tag 驱动，版本、CHANGELOG 和 Release Notes 必须一致。",
        "生产环境需要审批和回归验证。",
    ], header_color=Palette.primary)
    add_footer(slide, "发布不是把代码合进 main，而是把已验证的变更推到正确环境。")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(slide, "新员工上手清单", "ONB")
    checklist = [
        "已阅读流程总览。",
        "已理解角色分工。",
        "已理解 docs 分支和 feature 分支。",
        "已理解 Issue 关联规则。",
        "已理解本地验证方式。",
        "已理解 release 流程。",
    ]
    for idx, item in enumerate(checklist):
        y = 1.72 + idx * 0.62
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(y), Inches(0.34), Inches(0.34))
        box.fill.solid()
        box.fill.fore_color.rgb = Palette.green if idx < 3 else Palette.primary
        box.line.fill.background()
        add_textbox(slide, 1.42, y - 0.01, 10.8, 0.34, item, size=16, color=Palette.text)
    add_card(slide, 0.95, 5.75, 11.4, 0.72, "下一步", ["跟着一个真实需求，从 Issue 跑到 docs、feature、PR、发布。"], header_color=Palette.accent)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill_bg(slide)
    add_textbox(slide, 0.9, 2.1, 11.6, 0.7, "Q&A", size=36, color=Palette.primary, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 0.9, 3.0, 11.6, 0.4, "现在可以开始参与一个真实项目了", size=18, color=Palette.muted, align=PP_ALIGN.CENTER)
    add_card(slide, 2.0, 4.0, 9.25, 1.1, "总结", ["需求清晰、角色清晰、分支清晰、门禁清晰，团队协作就会稳定很多。"], header_color=Palette.green)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)


if __name__ == "__main__":
    build_deck(Path("docs/reports/health_platform_workflow_roles_ppt.pptx"))