"""Generate an executive-style PPTX deck for AI-enabled software delivery."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


class Palette:
    bg = RGBColor(247, 249, 252)
    text = RGBColor(33, 41, 54)
    muted = RGBColor(92, 106, 127)
    primary = RGBColor(24, 88, 124)
    accent = RGBColor(237, 115, 46)
    success = RGBColor(46, 125, 50)
    danger = RGBColor(190, 60, 60)
    card = RGBColor(255, 255, 255)
    border = RGBColor(220, 226, 233)
    blue2 = RGBColor(67, 128, 180)


def set_font(run, size_pt: int, bold: bool = False, color: RGBColor | None = None, name: str = "Microsoft YaHei"):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def fill_background(slide, color: RGBColor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header(slide, title: str, tag: str | None = None):
    fill_background(slide, Palette.bg)
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(0.45), Inches(12.2), Inches(0.9))
    box.fill.solid()
    box.fill.fore_color.rgb = Palette.card
    box.line.color.rgb = Palette.border

    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.63), Inches(10.4), Inches(0.55))
    tf = tb.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_font(r, 24, bold=True, color=Palette.primary)

    if tag:
        chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.35), Inches(0.58), Inches(1.15), Inches(0.42))
        chip.fill.solid()
        chip.fill.fore_color.rgb = Palette.primary
        chip.line.fill.background()
        ctf = chip.text_frame
        ctf.clear()
        cp = ctf.paragraphs[0]
        cp.alignment = PP_ALIGN.CENTER
        cr = cp.add_run()
        cr.text = tag
        set_font(cr, 11, bold=True, color=RGBColor(255, 255, 255))


def add_title_slide(slide, title: str, subtitle: str, footer: str):
    fill_background(slide, Palette.bg)
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.25)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = Palette.primary
    slide.shapes[-1].line.fill.background()

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.95), Inches(1.25), Inches(2.1), Inches(0.16))
    bar.fill.solid()
    bar.fill.fore_color.rgb = Palette.accent
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.95), Inches(1.55), Inches(10.7), Inches(1.2))
    tf = tb.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_font(r, 30, bold=True, color=Palette.primary)

    sub = slide.shapes.add_textbox(Inches(0.95), Inches(2.75), Inches(11.0), Inches(0.75))
    tf = sub.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = subtitle
    set_font(r, 18, color=Palette.muted)

    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(3.65), Inches(5.2), Inches(0.6))
    badge.fill.solid()
    badge.fill.fore_color.rgb = Palette.card
    badge.line.color.rgb = Palette.border
    tf = badge.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = footer
    set_font(r, 14, bold=True, color=Palette.primary)

    for idx, x in enumerate([8.6, 9.4, 10.2, 11.0, 11.8]):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(3.95), Inches(0.22), Inches(0.22))
        dot.fill.solid()
        dot.fill.fore_color.rgb = Palette.accent if idx % 2 == 0 else Palette.success
        dot.line.fill.background()


def add_card(slide, left: float, top: float, width: float, height: float, title: str, items: list[str], title_color: RGBColor | None = None):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = Palette.card
    card.line.color.rgb = Palette.border

    tb = slide.shapes.add_textbox(Inches(left + 0.16), Inches(top + 0.1), Inches(width - 0.32), Inches(0.34))
    tf = tb.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_font(r, 16, bold=True, color=title_color or Palette.primary)

    body = slide.shapes.add_textbox(Inches(left + 0.16), Inches(top + 0.48), Inches(width - 0.32), Inches(height - 0.58))
    tf = body.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, text in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(5)
        p.text = text
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(15)
        p.font.color.rgb = Palette.text


def add_flow_boxes(slide, boxes: list[tuple[str, str, RGBColor]], y: float = 2.05):
    x0 = 0.7
    box_w = 1.95
    gap = 0.15
    for idx, (title, desc, color) in enumerate(boxes):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x0 + idx * (box_w + gap)), Inches(y), Inches(box_w), Inches(1.15))
        box.fill.solid()
        box.fill.fore_color.rgb = color
        box.line.fill.background()

        tb1 = slide.shapes.add_textbox(Inches(x0 + 0.12 + idx * (box_w + gap)), Inches(y + 0.14), Inches(box_w - 0.24), Inches(0.3))
        tf = tb1.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = title
        set_font(r, 13, bold=True, color=RGBColor(255, 255, 255))

        tb2 = slide.shapes.add_textbox(Inches(x0 + 0.12 + idx * (box_w + gap)), Inches(y + 0.5), Inches(box_w - 0.24), Inches(0.44))
        tf = tb2.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = desc
        set_font(r, 10, color=RGBColor(255, 255, 255))

        if idx < len(boxes) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x0 + (idx + 1) * (box_w + gap) - 0.1), Inches(y + 0.38), Inches(0.28), Inches(0.3))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = Palette.accent
            arrow.line.fill.background()


def add_three_column_matrix(slide, rows: list[tuple[str, str, str]]):
    y0 = 1.7
    heights = [1.2, 1.2, 1.2, 1.2]
    x_positions = [0.7, 4.7, 8.7]
    widths = [3.6, 3.6, 3.6]
    for idx, (title, mid, right) in enumerate(rows):
        title_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x_positions[0]), Inches(y0 + idx * 1.25), Inches(widths[0]), Inches(heights[idx]))
        title_box.fill.solid(); title_box.fill.fore_color.rgb = Palette.primary; title_box.line.fill.background()
        tb = slide.shapes.add_textbox(Inches(x_positions[0] + 0.12), Inches(y0 + idx * 1.25 + 0.1), Inches(widths[0] - 0.24), Inches(heights[idx] - 0.2))
        tf = tb.text_frame; tf.clear(); p = tf.paragraphs[0]; r = p.add_run(); r.text = title; set_font(r, 12, bold=True, color=RGBColor(255,255,255))

        mid_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x_positions[1]), Inches(y0 + idx * 1.25), Inches(widths[1]), Inches(heights[idx]))
        mid_box.fill.solid(); mid_box.fill.fore_color.rgb = Palette.card; mid_box.line.color.rgb = Palette.border
        tb = slide.shapes.add_textbox(Inches(x_positions[1] + 0.12), Inches(y0 + idx * 1.25 + 0.1), Inches(widths[1] - 0.24), Inches(heights[idx] - 0.2))
        tf = tb.text_frame; tf.clear(); p = tf.paragraphs[0]; r = p.add_run(); r.text = mid; set_font(r, 11, color=Palette.text)

        right_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x_positions[2]), Inches(y0 + idx * 1.25), Inches(widths[2]), Inches(heights[idx]))
        right_box.fill.solid(); right_box.fill.fore_color.rgb = Palette.card; right_box.line.color.rgb = Palette.border
        tb = slide.shapes.add_textbox(Inches(x_positions[2] + 0.12), Inches(y0 + idx * 1.25 + 0.1), Inches(widths[2] - 0.24), Inches(heights[idx] - 0.2))
        tf = tb.text_frame; tf.clear(); p = tf.paragraphs[0]; r = p.add_run(); r.text = right; set_font(r, 11, color=Palette.text)


def build_deck(out_path: Path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    today = datetime.now().strftime("%Y-%m-%d")
    version = (Path(__file__).resolve().parents[1] / "VERSION").read_text(encoding="utf-8").strip()

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_slide(
        slide,
        "AI 赋能的软件研发流程",
        "面向高层汇报：从需求到上线，AI 如何重塑协作与价值",
        f"{today} · Version {version}",
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "为什么这件事值得被重视", "WHY")
    add_card(slide, 0.8, 1.55, 5.85, 4.95, "业务背景", [
        "软件研发不再只是“写代码”，而是持续交付价值。",
        "AI 能把需求澄清、设计、编码、测试和交付加速。",
        "高层关注的不只是效率，而是质量、速度和可控性。",
    ], Palette.primary)
    add_card(slide, 6.95, 1.55, 5.55, 4.95, "汇报要点", [
        "AI 不是替代团队，而是重构流程。",
        "人负责判断、责任和业务决策。",
        "AI 负责加速信息加工、生成草案与重复性工作。",
    ], Palette.accent)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "通用的软件研发流程", "FLOW")
    add_flow_boxes(slide, [
        ("需求", "目标\n范围\n价值", Palette.primary),
        ("设计", "架构\n方案\n计划", Palette.blue2),
        ("开发", "实现\n集成\n交付", Palette.accent),
        ("测试", "验证\n回归\n质量", Palette.success),
        ("发布", "上线\n观察\n优化", Palette.primary),
    ], y=2.2)
    add_card(slide, 0.8, 4.15, 11.75, 1.8, "核心结论", [
        "整个流程围绕“需求→设计→实现→验证→发布→运营”展开，AI 可以在每一步嵌入。"
    ], Palette.success)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "AI 在每个阶段扮演什么角色", "AI BY STAGE")
    add_three_column_matrix(slide, [
        ("需求", "整理背景、生成问题清单、辅助验收标准", "人工确认业务价值与优先级"),
        ("设计", "生成方案草稿、结构化设计建议、风险提示", "人工最终裁决与架构把关"),
        ("开发", "生成代码骨架、文档、测试样例、重构建议", "人工负责实现细节与责任边界"),
        ("测试与发布", "自动生成测试、回归检查、发布说明", "人工负责审批、监控和发布决策"),
    ])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "人工、AI 与协作的分工模型", "MODEL")
    add_card(slide, 0.8, 1.55, 3.8, 4.95, "人工负责", [
        "业务目标与优先级",
        "最终决策与责任",
        "合规、安全、风险判断",
        "复杂场景的经验性判断",
    ], Palette.primary)
    add_card(slide, 4.7, 1.55, 3.8, 4.95, "AI 负责", [
        "信息汇总与结构化",
        "草稿生成与建议输出",
        "重复性编码和测试",
        "快速分析与知识检索",
    ], Palette.blue2)
    add_card(slide, 8.6, 1.55, 3.8, 4.95, "协作价值", [
        "人做判断，AI做加速",
        "让团队把更多时间放在价值创造",
        "缩短从想法到可交付成果的周期",
    ], Palette.success)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "AI 能带来的价值", "VALUE")
    add_card(slide, 0.8, 1.55, 2.8, 4.95, "速度", [
        "需求转成文档和方案更快",
        "编码和测试更快",
        "周转时间明显缩短",
    ], Palette.primary)
    add_card(slide, 3.75, 1.55, 2.8, 4.95, "质量", [
        "规范化输出更一致",
        "测试覆盖更早进入流程",
        "错误更容易被发现和修补",
    ], Palette.blue2)
    add_card(slide, 6.7, 1.55, 2.8, 4.95, "协作", [
        "跨角色沟通更顺畅",
        "知识沉淀更容易复用",
        "新成员更快速上手",
    ], Palette.accent)
    add_card(slide, 9.65, 1.55, 2.8, 4.95, "治理", [
        "可追踪、可审查、可回退",
        "AI 输出与人工审核并行",
        "风险可控且可持续优化",
    ], Palette.success)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "在本项目中的落地方式", "PRACTICE")
    add_card(slide, 0.8, 1.55, 5.9, 4.95, "流程落地", [
        "需求阶段用 AI 生成问题清单与文档草稿。",
        "设计与计划阶段用 AI 辅助结构化输出。",
        "开发阶段用 AI 加速代码、测试与文档。",
        "发布前仍然需要人工确认和质量门禁。",
    ], Palette.primary)
    add_card(slide, 6.95, 1.55, 5.55, 4.95, "组织价值", [
        "让团队把精力从重复劳动转向判断与创新。",
        "形成可复制的研发标准与协作方式。",
        "为高层提供更清晰的交付节奏与风险可控性。",
    ], Palette.accent)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "结论与建议", "END")
    add_card(slide, 1.0, 1.9, 11.3, 3.3, "一句话总结", [
        "AI 的价值不在替代人，而在帮助团队把“想法”更快、更加一致地变成可交付的软件成果。"
    ], Palette.success)
    add_card(slide, 1.0, 5.4, 11.3, 1.05, "建议", [
        "把 AI 作为研发流程中的“协作伙伴”来设计，而不是“额外工具”来堆叠。"
    ], Palette.primary)

    prs.save(out_path)


def main():
    repo_root = Path(__file__).resolve().parents[1]
    version = (repo_root / "VERSION").read_text(encoding="utf-8").strip()
    today = datetime.now().strftime("%Y-%m-%d")
    out_path = repo_root / "docs" / "reports" / f"health_platform_ai_workflow_{today}_v{version}.pptx"
    build_deck(out_path)
    print(out_path)


if __name__ == "__main__":
    main()
