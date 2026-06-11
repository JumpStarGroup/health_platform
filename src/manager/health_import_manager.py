"""Manager for health record import (preview + commit)."""
import csv
import json
import uuid
from datetime import datetime, timezone, timedelta
from io import StringIO, BytesIO
from typing import Dict, List, Optional, Tuple

from dateutil import parser as dt_parser, tz
try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - optional dependency
    load_workbook = None
from sqlalchemy import and_, select

from ..extensions import db
from ..models import (
    HealthRecord,
    RecordSubject,
    Member,
    HealthImportSession,
    HealthImportAudit,
)
from ..timeutil import UTC
from .member_manager import MemberManager
from .health_manager import DuplicateHealthRecordError, HealthManager


class ImportValidationError(ValueError):
    """Raised when import preview or commit fails due to user-correctable issues."""


class HealthImportManager:
    MAX_ROWS = 500
    MAX_FILE_BYTES = 5 * 1024 * 1024

    REQUIRED_FIELDS = ["member_name", "timestamp", "systolic", "diastolic"]
    OPTIONAL_FIELDS = ["heart_rate", "tags", "note"]

    # Known header aliases (normalized to lower/stripped)
    HEADER_ALIASES = {
        "member_name": {
            "member name",
            "member",
            "name",
            "full name",
            "成员",
            "成员名称",
            "家庭成员",
            "姓名",
        },
        "timestamp": {
            "timestamp",
            "time",
            "datetime",
            "测量时间",
            "时间",
            "日期",
            "日期时间",
        },
        "systolic": {
            "systolic",
            "systolic pressure",
            "systolic_pressure",
            "收缩压",
            "高压",
        },
        "diastolic": {
            "diastolic",
            "diastolic pressure",
            "diastolic_pressure",
            "舒张压",
            "低压",
        },
        "heart_rate": {
            "heart rate",
            "heartrate",
            "heart_rate",
            "心率",
            "脉搏",
        },
        "tags": {
            "tags",
            "tag",
            "标签",
        },
        "note": {
            "note",
            "备注",
            "说明",
            "comment",
        },
    }

    def __init__(self):
        self.member_mgr = MemberManager()
        self.health_mgr = HealthManager()

    # ---------- Public entrypoints ----------
    def preview(self, user_id: int, file_storage, mapping: Optional[Dict[str, str]] = None) -> Dict:
        if not file_storage:
            raise ImportValidationError("file is required")
        filename = file_storage.filename or ""
        if not filename.lower().endswith((".csv", ".xlsx")):
            raise ImportValidationError("unsupported file type")

        file_bytes = file_storage.read()
        if len(file_bytes) > self.MAX_FILE_BYTES:
            raise ImportValidationError("file exceeds 5MB limit")

        columns, rows = self._parse_file(filename, file_bytes)
        if len(rows) > self.MAX_ROWS:
            raise ImportValidationError("row limit exceeded (500)")

        mapping_resolved, missing_required = self._resolve_mapping(columns, mapping or {})
        if missing_required:
            return {
                "session_id": None,
                "columns": columns,
                "mapping": mapping_resolved,
                "missing_required_fields": list(missing_required),
                "preview": {"total_rows": len(rows), "valid_count": 0, "error_count": len(rows), "skipped_count": 0},
                "unknown_members": [],
                "rows": [],
            }

        preview_rows, unknown_members, skipped_rows = self._build_preview_rows(user_id, rows, mapping_resolved)

        session_id = str(uuid.uuid4())
        session = HealthImportSession(
            id=session_id,
            user_id=user_id,
            original_filename=filename,
            file_size_bytes=len(file_bytes),
            mapping_json=json.dumps(mapping_resolved, ensure_ascii=False),
            normalized_rows_json=json.dumps(preview_rows, ensure_ascii=False),
            preview_errors_json=json.dumps([r for r in preview_rows if r["status"] == "error"], ensure_ascii=False),
            preview_skipped_json=json.dumps(skipped_rows, ensure_ascii=False),
            unknown_members_json=json.dumps(sorted(list(unknown_members)), ensure_ascii=False),
            status="previewed",
            created_at=self._utc_now_naive(),
            expires_at=self._utc_now_naive() + timedelta(hours=24),
        )
        db.session.add(session)
        db.session.commit()

        # Get existing members for mapping UI
        existing_members = self.member_mgr.list_members(user_id)
        existing_members_list = [{"id": m.id, "name": m.full_name} for m in existing_members]

        counts = self._count_by_status(preview_rows)
        return {
            "session_id": session_id,
            "columns": columns,
            "mapping": mapping_resolved,
            "missing_required_fields": [],
            "preview": {
                "total_rows": len(rows),
                "valid_count": counts.get("valid", 0),
                "error_count": counts.get("error", 0),
                "skipped_count": counts.get("skipped_duplicate", 0),
            },
            "unknown_members": sorted(list(unknown_members)),
            "existing_members": existing_members_list,
            "rows": preview_rows,
        }

    def commit(
        self,
        user_id: int,
        session_id: str,
        approved_new_members: Optional[List[str]] = None,
        member_mappings: Optional[Dict[str, int]] = None,
    ) -> Dict:
        """
        Commit import session.
        
        Args:
            approved_new_members: List of new member names to create
            member_mappings: Dict mapping Excel member names to existing member IDs
        """
        approved_new_members = approved_new_members or []
        member_mappings = member_mappings or {}
        session = HealthImportSession.query.filter_by(id=session_id, user_id=user_id).first()
        if not session:
            raise ImportValidationError("session not found")
        if session.status != "previewed":
            raise ImportValidationError("session already committed or expired")
        expires_at = self._ensure_naive_utc(session.expires_at) if session.expires_at else None
        if expires_at and expires_at < self._utc_now_naive():
            raise ImportValidationError("session expired")

        preview_rows = json.loads(session.normalized_rows_json or "[]")
        mapping = json.loads(session.mapping_json or "{}")
        unknown_members = {m.strip(): m.strip() for m in json.loads(session.unknown_members_json or "[]") if m.strip()}
        approved_norm = {self._normalize_name(n): n for n in approved_new_members if n}

        # Build member map from existing members
        member_map = self._load_member_map(user_id)
        
        # Apply user-provided mappings (Excel name -> existing member_id)
        for excel_name, member_id in member_mappings.items():
            norm_name = self._normalize_name(excel_name)
            if member_id:
                member_map[norm_name] = member_id
        
        # Create approved new members
        created_members_count = 0
        for norm_name, display_name in unknown_members.items():
            if norm_name in approved_norm and norm_name not in member_map:
                created = self._create_member(user_id, display_name)
                member_map[self._normalize_name(created.full_name)] = created.id
                created_members_count += 1

        # Re-evaluate rows for commit
        commit_rows = []
        for r in preview_rows:
            # Skip rows that were hard errors in preview
            if r.get("status") == "error":
                commit_rows.append(r)
                continue
            member_name = r.get("data", {}).get("member_name")
            norm_member = self._normalize_name(member_name or "")
            member_id = member_map.get(norm_member)
            if not member_id:
                r["status"] = "error"
                r.setdefault("messages", []).append("member not approved or not found")
                commit_rows.append(r)
                continue
            r["member_id"] = member_id
            commit_rows.append(r)

        # Duplicate detection at commit time
        self._mark_duplicates(user_id, commit_rows)

        success_count = 0
        skipped_count = 0
        error_count = 0

        for r in commit_rows:
            status = r.get("status")
            if status == "error":
                error_count += 1
                continue
            if status == "skipped_duplicate":
                skipped_count += 1
                continue
            data = r.get("data", {})
            try:
                rec = self.health_mgr.create(
                    user_id=user_id,
                    systolic=data.get("systolic"),
                    diastolic=data.get("diastolic"),
                    heart_rate=data.get("heart_rate"),
                    timestamp=self._parse_timestamp_to_utc(data.get("timestamp")),
                    tags=data.get("tags") or [],
                    note=data.get("note"),
                    subject_member_id=r.get("member_id"),
                    treat_unmapped_as_self=False,
                )
                # Map record to subject
                hh = self.member_mgr.ensure_default_household(user_id)
                rs = RecordSubject(
                    record_id=rec.id,
                    household_id=hh.id,
                    member_id=r.get("member_id"),
                    created_by_user_id=user_id,
                )
                db.session.add(rs)
                success_count += 1
            except DuplicateHealthRecordError:
                r["status"] = "skipped_duplicate"
                skipped_count += 1
            except Exception as exc:
                r["status"] = "error"
                r.setdefault("messages", []).append(str(exc))
                error_count += 1
        db.session.commit()

        total_rows = len(commit_rows)
        # Audit
        audit = HealthImportAudit(
            user_id=user_id,
            session_id=session_id,
            original_filename=session.original_filename,
            file_size_bytes=session.file_size_bytes,
            file_sha256=session.file_sha256,
            total_rows=total_rows,
            success_count=success_count,
            error_count=error_count,
            skipped_count=skipped_count,
            created_members_count=created_members_count,
            created_at=self._utc_now_naive(),
        )
        db.session.add(audit)
        session.status = "committed"
        db.session.commit()

        errors_payload = [r for r in commit_rows if r.get("status") == "error"]
        skipped_payload = [r for r in commit_rows if r.get("status") == "skipped_duplicate"]

        return {
            "summary": {
                "total_rows": total_rows,
                "success_count": success_count,
                "error_count": error_count,
                "skipped_count": skipped_count,
            },
            "report": {
                "errors": errors_payload,
                "skipped": skipped_payload,
            },
        }

    # ---------- Helpers ----------
    def _parse_file(self, filename: str, file_bytes: bytes) -> Tuple[List[str], List[Dict[str, str]]]:
        if filename.lower().endswith(".csv"):
            text = file_bytes.decode("utf-8-sig")
            reader = csv.reader(StringIO(text))
            rows = list(reader)
        else:
            if load_workbook is None:
                raise ImportValidationError("excel support requires openpyxl")
            wb = load_workbook(BytesIO(file_bytes), read_only=True)
            ws = wb.active
            # values_only=True already returns values, not cell objects
            rows = [[cell if cell is not None else "" for cell in row] for row in ws.iter_rows(values_only=True)]

        if not rows:
            raise ImportValidationError("file is empty")
        columns = [str(c).strip() for c in rows[0]]
        data_rows = rows[1:]
        normalized_rows = []
        for idx, r in enumerate(data_rows, start=2):
            row_dict = {}
            for c_idx, col in enumerate(columns):
                row_dict[col] = r[c_idx] if c_idx < len(r) else ""
            normalized_rows.append({"row": idx, "raw": row_dict})
        return columns, normalized_rows

    def _resolve_mapping(self, columns: List[str], mapping: Dict[str, str]) -> Tuple[Dict[str, str], set]:
        col_norm_map = {self._normalize_header(c): c for c in columns}
        resolved = {}
        # apply provided mapping first
        for field, col in (mapping or {}).items():
            if col in columns:
                resolved[field] = col

        for field in self.REQUIRED_FIELDS + self.OPTIONAL_FIELDS:
            if field in resolved:
                continue
            for alias in self.HEADER_ALIASES.get(field, set()):
                if alias in col_norm_map:
                    resolved[field] = col_norm_map[alias]
                    break
        missing = {f for f in self.REQUIRED_FIELDS if f not in resolved}
        return resolved, missing

    def _build_preview_rows(self, user_id: int, rows: List[Dict], mapping: Dict[str, str]):
        member_map = self._load_member_map(user_id)
        unknown_members = set()
        preview_rows = []

        # Preload columns for fast access
        for row in rows:
            row_num = row.get("row")
            raw = row.get("raw", {})
            data = {}
            messages = []
            status = "valid"

            # member
            member_name = raw.get(mapping.get("member_name"), "") if mapping.get("member_name") else ""
            member_name = self._normalize_display(member_name)
            if not member_name:
                status = "error"
                messages.append("member_name is required")
            data["member_name"] = member_name

            # timestamp
            ts_raw = raw.get(mapping.get("timestamp"), "") if mapping.get("timestamp") else ""
            try:
                ts = self._parse_timestamp_to_utc(ts_raw)
                data["timestamp"] = ts.isoformat().replace("+00:00", "Z")
            except Exception:
                status = "error"
                messages.append("invalid timestamp")

            # systolic
            try:
                data["systolic"] = self._parse_int_field(raw, mapping, "systolic", required=True)
            except ValueError as exc:
                status = "error"
                messages.append(str(exc))

            # diastolic
            try:
                data["diastolic"] = self._parse_int_field(raw, mapping, "diastolic", required=True)
            except ValueError as exc:
                status = "error"
                messages.append(str(exc))

            # heart rate
            try:
                data["heart_rate"] = self._parse_int_field(raw, mapping, "heart_rate", required=False)
            except ValueError as exc:
                status = "error"
                messages.append(str(exc))

            # relation check
            if data.get("systolic") is not None and data.get("diastolic") is not None:
                if data["systolic"] <= data["diastolic"]:
                    status = "error"
                    messages.append("systolic must be greater than diastolic")

            # tags
            tags_raw = raw.get(mapping.get("tags"), "") if mapping.get("tags") else ""
            data["tags"] = self._split_tags(tags_raw)

            # note
            note_raw = raw.get(mapping.get("note"), "") if mapping.get("note") else ""
            if note_raw:
                note_str = str(note_raw)
                if len(note_str) > 500:
                    status = "error"
                    messages.append("note exceeds 500 characters")
                data["note"] = note_str
            else:
                data["note"] = None

            # member resolution
            norm_member = self._normalize_name(member_name)
            member_id = member_map.get(norm_member)
            if not member_id and member_name:
                unknown_members.add(member_name)
                status = "unknown_member"
                messages.append("member not found")

            preview_rows.append({
                "row": row_num,
                "data": data,
                "status": status,
                "messages": messages,
                **({"member_id": member_id} if member_id else {}),
            })

        # Duplicate detection (only for rows with resolved member_id and valid)
        self._mark_duplicates(user_id, preview_rows)
        skipped_rows = [r for r in preview_rows if r.get("status") == "skipped_duplicate"]
        return preview_rows, unknown_members, skipped_rows

    def _mark_duplicates(self, user_id: int, rows: List[Dict]):
        keys = []  # (member_id, minute_ts)
        seen = set()
        for r in rows:
            if r.get("status") not in {"valid", "unknown_member"}:
                continue
            member_id = r.get("member_id")
            if not member_id:
                continue
            ts_iso = r.get("data", {}).get("timestamp")
            if not ts_iso:
                continue
            ts = self._parse_timestamp_to_utc(ts_iso)
            minute = self._normalize_minute(ts)
            if (member_id, minute) in seen:
                r["status"] = "skipped_duplicate"
                r.setdefault("messages", []).append("duplicate in the same minute (file)")
            else:
                seen.add((member_id, minute))
                keys.append((member_id, minute))
            r["_minute_ts"] = minute

        if not keys:
            return
        member_ids = {k[0] for k in keys}
        minutes = {k[1] for k in keys}
        existing = db.session.query(RecordSubject.member_id, HealthRecord.timestamp).join(
            HealthRecord, HealthRecord.id == RecordSubject.record_id
        ).filter(
            HealthRecord.user_id == user_id,
            RecordSubject.member_id.in_(member_ids),
            HealthRecord.timestamp.in_(list(minutes)),
        ).all()
        existing_set = {(mid, ts) for mid, ts in existing}

        for r in rows:
            member_id = r.get("member_id")
            minute = r.get("_minute_ts")
            if member_id and minute and (member_id, minute) in existing_set:
                r["status"] = "skipped_duplicate"
                r.setdefault("messages", []).append("duplicate in the same minute")
            if "_minute_ts" in r:
                r.pop("_minute_ts", None)

    def _load_member_map(self, user_id: int) -> Dict[str, int]:
        members = self.member_mgr.list_members(user_id)
        result = {}
        for m in members:
            norm = self._normalize_name(m.full_name)
            result[norm] = m.id
            # Self aliases
            if norm == "self":
                for alias in ["自己", "本人"]:
                    result[self._normalize_name(alias)] = m.id
        return result

    def _create_member(self, user_id: int, name: str) -> Member:
        # extend MemberManager create_member to handle "本人"
        normalized = name.strip()
        if self._normalize_name(normalized) in {"self", "自己", "本人"}:
            normalized = "Self"
        return self.member_mgr.create_member(user_id, normalized, None, None, None, None)

    def _parse_int_field(self, raw: Dict, mapping: Dict[str, str], field: str, required: bool) -> Optional[int]:
        col = mapping.get(field)
        val = raw.get(col, "") if col else ""
        if val is None or str(val).strip() == "":
            if required:
                raise ValueError(f"{field} is required")
            return None
        try:
            intval = int(str(val).strip())
        except Exception:
            raise ValueError(f"{field} must be an integer")
        if field in {"systolic", "diastolic"}:
            if not (30 <= intval <= 250):
                raise ValueError(f"{field} must be between 30-250")
        if field == "heart_rate" and intval is not None:
            if not (30 <= intval <= 150):
                raise ValueError("heart_rate must be between 30-150")
        return intval

    def _split_tags(self, value) -> List[str]:
        if not value:
            return []
        if isinstance(value, list):
            return [str(v).strip() for v in value if str(v).strip()]
        parts = str(value).replace(";", ",").split(",")
        return [p.strip() for p in parts if p.strip()]

    def _parse_timestamp_to_utc(self, value) -> datetime:
        if isinstance(value, datetime):
            dt = value
        else:
            dt = dt_parser.parse(str(value))
        if dt.tzinfo is None:
            tz_bj = tz.gettz("Asia/Shanghai")
            dt = dt.replace(tzinfo=tz_bj)
        dt_utc = dt.astimezone(timezone.utc)
        return dt_utc.replace(tzinfo=None)

    def _normalize_minute(self, dt_value: datetime) -> datetime:
        return dt_value.replace(second=0, microsecond=0)

    def _normalize_header(self, name: str) -> str:
        return (name or "").strip().lower()

    def _normalize_name(self, name: str) -> str:
        return (name or "").strip().lower()

    def _normalize_display(self, value) -> str:
        if value is None:
            return ""
        return str(value).strip()

    def _count_by_status(self, rows: List[Dict]) -> Dict[str, int]:
        counts = {}
        for r in rows:
            st = r.get("status") or ""
            counts[st] = counts.get(st, 0) + 1
        return counts

    def _utc_now_naive(self) -> datetime:
        return datetime.now(timezone.utc).replace(tzinfo=None)

    def _ensure_naive_utc(self, dt_value: datetime) -> datetime:
        if dt_value is None:
            return None
        if dt_value.tzinfo is None:
            return dt_value
        return dt_value.astimezone(timezone.utc).replace(tzinfo=None)
